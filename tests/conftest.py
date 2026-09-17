import io

import pytest

from jobcu.paths import DATA_DIR_ENV


@pytest.fixture(autouse=True)
def temporary_data_dir(tmp_path, monkeypatch):
    """Every test uses a throwaway data folder, never the real one."""
    folder = tmp_path / "jobcu-data"
    monkeypatch.setenv(DATA_DIR_ENV, str(folder))
    return folder


def make_pdf(lines: list[str]) -> bytes:
    """A tiny one-page PDF with the given lines of text, built by hand for tests."""

    def escape(text: str) -> str:
        return text.replace("\\", "\\\\").replace("(", "\\(").replace(")", "\\)")

    content = "BT /F1 11 Tf 50 750 Td 14 TL " + " ".join(f"({escape(x)}) Tj T*" for x in lines)
    content += " ET"
    objects = [
        "<< /Type /Catalog /Pages 2 0 R >>",
        "<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        "<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] "
        "/Resources << /Font << /F1 4 0 R >> >> /Contents 5 0 R >>",
        "<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
        f"<< /Length {len(content)} >>\nstream\n{content}\nendstream",
    ]
    out = b"%PDF-1.4\n"
    offsets = []
    for number, body in enumerate(objects, start=1):
        offsets.append(len(out))
        out += f"{number} 0 obj\n{body}\nendobj\n".encode("latin-1")
    xref = len(out)
    out += f"xref\n0 {len(objects) + 1}\n0000000000 65535 f \n".encode()
    for offset in offsets:
        out += f"{offset:010d} 00000 n \n".encode()
    trailer = f"trailer\n<< /Size {len(objects) + 1} /Root 1 0 R >>\nstartxref\n{xref}\n%%EOF\n"
    out += trailer.encode()
    return out


def make_docx(paragraphs: list[str], table_rows: list[list[str]] | None = None) -> bytes:
    import docx

    document = docx.Document()
    for text in paragraphs:
        document.add_paragraph(text)
    if table_rows:
        table = document.add_table(rows=len(table_rows), cols=len(table_rows[0]))
        for r, row in enumerate(table_rows):
            for c, value in enumerate(row):
                table.cell(r, c).text = value
    buffer = io.BytesIO()
    document.save(buffer)
    return buffer.getvalue()


FAKE_CV_LINES = [
    "Alex Example",
    "Hardware engineer with five years of experience in embedded systems.",
    "Skills: schematic design, PCB layout, C programming, oscilloscope measurements.",
    "Languages: English (fluent), French (basic).",
]
